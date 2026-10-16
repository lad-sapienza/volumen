"""Renderer DOCX (python-docx): consuma i blocchi del modello."""
import io
import re
import struct
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..config import heading_style, style_cfg
from .. import model
from ..model import INLINE_MD

try:
    from PIL import Image
except ImportError:  # le dimensioni immagine saranno solo a larghezza massima
    Image = None

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def apply_style(style, sc):
    f = style.font
    f.name = sc["font"]
    f.size = Pt(sc["size"])
    f.bold = bool(sc["bold"])
    f.italic = bool(sc["italic"])
    f.color.rgb = RGBColor.from_string(sc["color"].upper())
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), sc["font"])
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    pf = style.paragraph_format
    # Altezza di riga ESATTA in pt (come Typst): stesso risultato nei due output.
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(sc["line_height"])
    pf.first_line_indent = Pt(sc["first_line_indent"])  # pt; 0 = nessun rientro; negativo = sporgente
    if sc["indent_left"]:
        pf.left_indent = Pt(sc["indent_left"])
    if sc["indent_right"]:
        pf.right_indent = Pt(sc["indent_right"])
    pf.space_before = Pt(sc["space_before"])
    pf.space_after = Pt(sc["space_after"])
    pf.alignment = ALIGN.get(sc["align"], WD_ALIGN_PARAGRAPH.LEFT)


def get_or_add_style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


def setup_document(cfg):
    doc = Document()
    sec = doc.sections[0]
    pg = cfg.get("page", {})
    if pg.get("size", "A4").upper() == "A4":
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    else:  # Letter
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.top_margin = Cm(pg.get("margin_top_cm", 2.5))
    sec.bottom_margin = Cm(pg.get("margin_bottom_cm", 2.5))
    if "margin_interior" in pg or "margin_exterior" in pg:
        # margini speculari (rilegatura): left = interno, right = esterno
        sec.left_margin = Cm(pg.get("margin_interior", pg.get("margin_left_cm", 2.5)))
        sec.right_margin = Cm(pg.get("margin_exterior", pg.get("margin_right_cm", 2.5)))
        doc.settings.element.find(qn("w:defaultTabStop")).addprevious(OxmlElement("w:mirrorMargins"))
    else:
        sec.left_margin = Cm(pg.get("margin_left_cm", 2.5))
        sec.right_margin = Cm(pg.get("margin_right_cm", 2.5))

    # sillabazione automatica: spenta di default (typography.hyphenate), nessuna parola spezzata a fine riga
    hyph = OxmlElement("w:autoHyphenation")
    hyph.set(qn("w:val"), "true" if cfg.get("typography", {}).get("hyphenate", False) else "false")
    doc.settings.element.find(qn("w:defaultTabStop")).addnext(hyph)   # posizione prevista dallo schema
    apply_style(doc.styles["Normal"], style_cfg(cfg, "normal"))
    for n in range(1, 6):
        apply_style(doc.styles[f"Heading {n}"], style_cfg(cfg, heading_style(cfg, n)))
        if n >= 2:
            doc.styles[f"Heading {n}"].paragraph_format.keep_with_next = True
    apply_style(doc.styles["Caption"], style_cfg(cfg, "caption"))
    for custom, key in (("Volumen Table Header", "table_header"),
                        ("Volumen Table Body", "table_body"),
                        ("Volumen Quote", "quote"), ("Volumen List", "list"), ("Volumen Bibliography", "bibliography"),
                        ("Volumen Chapter Author", "chapter_author"),
                        ("Footnote Text", "footnote")):
        st = get_or_add_style(doc, custom)
        st.base_style = doc.styles["Normal"]
        apply_style(st, style_cfg(cfg, key, fallback="normal"))

    fr = get_or_add_style(doc, "Footnote Reference", WD_STYLE_TYPE.CHARACTER)
    fr.font.superscript = True
    usable = (sec.page_width - sec.left_margin - sec.right_margin)
    usable_h = (sec.page_height - sec.top_margin - sec.bottom_margin)
    return doc, usable, usable_h


def apply_run_font(run, sc):
    f = run.font
    f.name = sc["font"]
    f.size = Pt(sc["size"])
    f.bold = bool(sc["bold"])
    f.italic = bool(sc["italic"])
    f.color.rgb = RGBColor.from_string(sc["color"].upper())
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), sc["font"])


def add_page_field(paragraph, sc):
    run = paragraph.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)
    apply_run_font(run, sc)


def set_page_number_start(section, start):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:start"), str(start))


def add_running_header(doc, section, cfg, usable, title=None, first=True, left_title=None):
    """Intestazione corrente, specchiata fra pagine dispari e pari: il titolo
    (da config) verso l'interno (rilegatura), il numero di pagina verso
    l'esterno, con una riga sotto. Dispari: titolo a sinistra, numero a
    destra. Pari: numero a sinistra, a destra left_title (autore) se dato, altrimenti il titolo."""
    hcfg = cfg.get("header", {})
    if title is None:
        title = hcfg.get("title", "")
    show_pn = cfg.get("page", {}).get("page_numbers", True)
    rule = hcfg.get("rule", {})
    title_sc = style_cfg(cfg, "header_title")
    pn_sc = style_cfg(cfg, "header_pagenum")

    doc.settings.odd_and_even_pages_header_footer = True
    section.header.is_linked_to_previous = False
    section.even_page_header.is_linked_to_previous = False
    if hcfg.get("skip_chapter_first_page", True):
        # prima pagina della sezione (= del capitolo) senza intestazione: la first-page header e' vuota
        section.different_first_page_header_footer = True
        section.first_page_header.is_linked_to_previous = False
    if first:
        set_page_number_start(section, hcfg.get("start_page", 1))
    else:   # la numerazione prosegue: la sezione clonata non deve ripartire da start
        pg = section._sectPr.find(qn("w:pgNumType"))
        if pg is not None and pg.get(qn("w:start")) is not None:
            del pg.attrib[qn("w:start")]

    def add_rule(paragraph):
        if rule.get("enabled", True):
            pPr = paragraph._p.get_or_add_pPr()
            add_bottom_border(pPr, rule.get("size_eighths_pt", 4),
                               rule.get("color", "000000"), rule.get("space_pt", 2))

    p_odd = section.header.paragraphs[0]
    p_odd.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    apply_run_font(p_odd.add_run(title), title_sc)
    p_odd.add_run().add_tab()
    if show_pn:
        add_page_field(p_odd, pn_sc)
    add_rule(p_odd)

    p_even = section.even_page_header.paragraphs[0]
    p_even.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    if show_pn:
        add_page_field(p_even, pn_sc)
    p_even.add_run().add_tab()
    apply_run_font(p_even.add_run(title if left_title is None else left_title), title_sc)
    add_rule(p_even)


CT_FOOTNOTES = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_XML = (
    f'<w:footnotes xmlns:w="{W_NS}">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
    'w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
    'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '</w:footnotes>')


def footnotes_part(doc_part):
    """Parte /word/footnotes.xml del documento (creata al primo uso: python-docx non la gestisce)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import XmlPart
    from docx.oxml import parse_xml
    for rel in doc_part.rels.values():
        if rel.reltype == RT.FOOTNOTES:
            return rel.target_part
    part = XmlPart(PackURI("/word/footnotes.xml"), CT_FOOTNOTES, parse_xml(FOOTNOTES_XML),
                   doc_part.package)
    doc_part.relate_to(part, RT.FOOTNOTES)
    return part


def add_footnote(par, note_text):
    """Vera nota a pie di pagina Word: richiamo (apice) nel paragrafo + testo in footnotes.xml,
    con gli stili 'Footnote Text' / 'Footnote Reference' (styles.footnote)."""
    from docx.text.paragraph import Paragraph
    root = footnotes_part(par.part).element
    ids = [int(f.get(qn("w:id"))) for f in root.findall(qn("w:footnote"))]
    nid = max(ids + [0]) + 1

    fn = OxmlElement("w:footnote")
    fn.set(qn("w:id"), str(nid))
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), "FootnoteText")
    pPr.append(ps)
    p.append(pPr)
    fn.append(p)
    root.append(fn)
    note_par = Paragraph(p, par._parent)
    ref = note_par.add_run()
    ref.style = par.part.document.styles["Footnote Reference"]
    ref._r.append(OxmlElement("w:footnoteRef"))
    note_par.add_run(" ")
    add_md_runs(note_par, note_text)

    run = par.add_run()
    run.style = par.part.document.styles["Footnote Reference"]
    fr = OxmlElement("w:footnoteReference")
    fr.set(qn("w:id"), str(nid))
    run._r.append(fr)


HIGHLIGHT = "FFEB3B"      # colore evidenziato (typography.highlight_color)
LINKS = {}                # stile dei link (config `links`: color, underline, bold); vuoto = come il testo


def add_chapter_link(par, anchor, text):
    """Collegamento interno al capitolo (segnalibro `anchor`): le run restano quelle normali,
    senza stile 'Hyperlink', cosi' la resa non cambia."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    par._p.append(link)
    n = len(par.runs)
    add_md_runs(par, text)
    for r in par.runs[n:]:          # le run create dopo il link vanno dentro il link
        if LINKS.get("color"):
            r.font.color.rgb = RGBColor.from_string(LINKS["color"].upper())
        if LINKS.get("underline"):
            r.font.underline = True
        if LINKS.get("bold"):
            r.bold = True
        link.append(r._r)


def add_bookmark(par, name, bid):
    """Segnalibro attorno al paragrafo (destinazione dei link tra capitoli)."""
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    for el in (start, end):
        el.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    pPr = par._p.pPr
    if pPr is not None:
        pPr.addnext(start)
    else:
        par._p.insert(0, start)
    par._p.append(end)


TRACK_ID = [9000]


def track_change(par, runs, deleted):
    """Avvolge i run in <w:ins> / <w:del> (modifica tracciata, autore e data da annotations)."""
    from datetime import datetime, timezone
    if not runs:
        return
    TRACK_ID[0] += 1
    wrap = OxmlElement("w:del" if deleted else "w:ins")
    wrap.set(qn("w:id"), str(TRACK_ID[0]))
    wrap.set(qn("w:author"), model.ANNOT["author"] or "Author")
    wrap.set(qn("w:date"), datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    runs[0]._r.addprevious(wrap)
    for r in runs:
        if deleted:
            for t in r._r.findall(qn("w:t")):
                t.tag = qn("w:delText")
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        wrap.append(r._r)


def add_styled_text(doc, st, usable_cm=None):
    """Paragrafi con stile personalizzato (`::: {attributi}`): font, corpo, peso, allineamento, spazi."""
    sc = st.style
    paras = st.text.split("\n\n")
    # width: i paragrafi si restringono con rientri destro/sinistro secondo l'allineamento del blocco
    left_in = right_in = 0.0
    wcm = model.length_cm(sc.get("width"), usable_cm or 0)
    if wcm and usable_cm and wcm < usable_cm:
        free = usable_cm - wcm
        if sc["align"] == "right":
            left_in = free
        elif sc["align"] == "center":
            left_in = right_in = free / 2
        else:
            right_in = free
    for i, text in enumerate(paras):
        p = doc.add_paragraph(style="Normal")
        pf = p.paragraph_format
        p.alignment = ALIGN.get(sc["align"], WD_ALIGN_PARAGRAPH.LEFT)
        pf.space_before = Pt(sc["space_before"] if i == 0 else 0)
        pf.space_after = Pt(sc["space_after"] if i == len(paras) - 1 else 0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(sc["line_height"])
        pf.first_line_indent = Pt(sc["first_line_indent"])
        pf.left_indent = Cm(left_in)
        pf.right_indent = Cm(right_in)
        n = len(p.runs)
        add_md_runs(p, text.lower() if sc.get("case") == "lower" else text)
        for r in p.runs[n:]:
            r.font.name = sc["font"]
            rfonts = r._r.get_or_add_rPr().find(qn("w:rFonts"))
            if rfonts is not None:
                for attr in ("w:cs", "w:eastAsia"):
                    rfonts.set(qn(attr), sc["font"])
            r.font.size = Pt(sc["size"])
            if sc["bold"]:
                r.bold = True
            if sc["italic"]:
                r.italic = True
            r.font.color.rgb = RGBColor.from_string(sc["color"].upper())
            if sc.get("case") == "upper":
                r.font.all_caps = True
            elif sc.get("case") == "smallcaps":
                r.font.small_caps = True
            if sc.get("tracking"):
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:val"), str(int(round(sc["tracking"] * 20))))     # ventesimi di punto
                r._r.get_or_add_rPr().append(sp)


def add_md_runs(par, text):
    # Le note vanno risolte per prime: possono contenere un link a capitolo, che altrimenti spezzerebbe
    # il testo fra NOTE_OPEN e NOTE_CLOSE (e il ramo dei link richiamerebbe se stesso all'infinito).
    if model.NOTE_OPEN in text:
        pos = 0
        for m in model.NOTE.finditer(text):
            add_md_runs(par, text[pos:m.start()])
            add_footnote(par, m.group(1))
            pos = m.end()
        rest = text[pos:]
        if model.NOTE_OPEN in rest:      # nota non chiusa (testo troncato): si toglie il marcatore
            rest = rest.replace(model.NOTE_OPEN, "")
        add_md_runs(par, rest)
        return
    if model.LINK_OPEN in text:
        pos = 0
        for m in model.CHLINK.finditer(text):
            add_md_runs(par, text[pos:m.start()])
            add_chapter_link(par, model.chapter_anchor(m.group(1)), m.group(2))
            pos = m.end()
        rest = text[pos:]
        if model.LINK_OPEN in rest:      # link non chiuso: si toglie il marcatore
            rest = rest.replace(model.LINK_OPEN, "")
        add_md_runs(par, rest)
        return
    for part in INLINE_MD.split(text):
        if not part:
            continue
        if model.critic_parts(part) is not None:
            text, comment = model.critic_parts(part)
            n = len(par.runs)
            if text:
                add_md_runs(par, text)
            if model.ANNOT["enabled"] and DOC is not None and par.part is DOC.part:
                runs = par.runs[n:] or [par.add_run("")]     # commento puntuale: ancora vuota
                DOC.add_comment(runs, text=comment, author=model.ANNOT["author"],
                                initials=model.ANNOT["initials"])
        elif model.critic_edit(part) is not None:
            old, new = model.critic_edit(part)
            if model.ANNOT["enabled"]:                   # revisione: vere modifiche tracciate di Word
                if old:
                    n = len(par.runs)
                    add_md_runs(par, old)
                    track_change(par, par.runs[n:], deleted=True)
                if new:
                    n = len(par.runs)
                    add_md_runs(par, new)
                    track_change(par, par.runs[n:], deleted=False)
            else:
                keep = model.accepted_edit(old, new)
                if keep:
                    add_md_runs(par, keep)
        elif model.mark_text(part) is not None:
            n = len(par.runs)
            add_md_runs(par, model.mark_text(part))      # dentro l'evidenziato: corsivo, grassetto...
            for r in par.runs[n:]:
                shd = OxmlElement("w:shd")               # sfondo a colore libero (piu' di highlight di Word)
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), HIGHLIGHT)
                r._r.get_or_add_rPr().append(shd)
        elif part.startswith("~~") and part.endswith("~~") and len(part) > 4:
            n = len(par.runs)
            add_md_runs(par, part[2:-2])
            for r in par.runs[n:]:
                r.font.strike = True
        elif part.startswith("^") and part.endswith("^") and len(part) > 2:
            par.add_run(part[1:-1]).font.superscript = True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2) or \
             (part.startswith("_") and part.endswith("_") and len(part) > 2):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def add_text_block(doc, text):
    """Testo del DB: paragrafi separati da riga vuota, a-capo singoli = line break."""
    if not text or not text.strip():
        return False
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = [ln.rstrip() for ln in block.split("\n") if ln.strip()]
        if not lines:
            continue
        p = doc.add_paragraph(style="Normal")
        for i, ln in enumerate(lines):
            if i:
                p.add_run().add_break()
            add_md_runs(p, ln.strip())
    return True


def set_cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, (sz, val) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(int(sz)))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)


def set_table_borders(table, rule_pt):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    sz = int(round(rule_pt * 8))
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cell_margins(table, top=20, bottom=20, left=50, right=50):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblPr.append(m)


def add_spacer(doc, pt):
    """Spazio verticale esatto (in punti), indipendente da cosa precede/segue:
    riga vuota a interlinea singola con testo a corpo minimo, così l'unico
    contributo visibile è lo space_after impostato."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pt)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().font.size = Pt(1)


def add_table(doc, header, rows, cfg, usable, ratios=None, aligns=None, empty=None, caption="",
              markup=True):
    """Tabella con intestazione ripetuta a ogni pagina; larghezze da ratios (o table.col_ratios),
    bordi/regole da table.*; celle con markup inline (**grassetto**, *corsivo*)."""
    tcfg = cfg.get("table", {})
    n = len(header)
    ratios = ratios or [1] * n
    tot = float(sum(ratios))
    widths = [int(usable * r / tot) for r in ratios]
    aligns = aligns or ["left"] * n
    if caption:
        cp = doc.add_paragraph(style="Caption")
        add_md_runs(cp, caption)
        cp.paragraph_format.keep_with_next = True

    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(widths) / 635)))   # EMU -> twips
    for i, col in enumerate(table.columns):
        col.width = widths[i]
    set_table_borders(table, tcfg.get("rule_pt", 0.75))
    set_cell_margins(table)

    def fill(cell, value, style):
        p = cell.paragraphs[0]
        p.style = doc.styles[style]
        if markup:
            add_md_runs(p, str(value))
        else:
            p.add_run(str(value))
        return p

    hdr = table.rows[0]
    set_repeat_header(hdr)
    set_cant_split(hdr)
    hrule = int(round(tcfg.get("header_rule_pt", 0.5) * 8))
    for i, name in enumerate(header):
        c = hdr.cells[i]
        c.width = widths[i]
        p = fill(c, name, "Volumen Table Header")
        if aligns[i] != "left":
            p.alignment = ALIGN.get(aligns[i], WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_borders(c, bottom=(hrule, "single"))
    for row in rows:
        r = table.add_row()
        set_cant_split(r)
        for i, val in enumerate(row):
            c = r.cells[i]
            c.width = widths[i]
            v = val if (val not in (None, "") or empty is None) else empty
            p = fill(c, v if v is not None else "", "Volumen Table Body")
            if aligns[i] != "left":
                p.alignment = ALIGN.get(aligns[i], WD_ALIGN_PARAGRAPH.LEFT)
    # se non ci sono righe, la tabella resta con solo l'intestazione ("sempre mostrata")
    space_after = tcfg.get("space_after_pt", 14)
    if space_after:
        add_spacer(doc, space_after)
    return table


def image_size(path):
    if Image is None:
        return None
    try:
        with Image.open(path) as im:
            return im.size
    except Exception:
        return None


def is_landscape(path):
    size = image_size(path)
    if not size or not size[0] or not size[1]:
        return True  # dimensioni non leggibili: trattata come orizzontale
    w, h = size
    return w >= h


def _exif_tiff_start(data):
    """Offset dell'header TIFF nel primo segmento APP1 'Exif' di un JPEG, None se
    assente. Scorre i marker: l'Exif non è sempre il primo segmento (i JPEG DJI
    hanno prima un APP0/JFIF)."""
    if data[0:2] != b"\xff\xd8":
        return None
    off = 2
    while off + 4 <= len(data) and data[off] == 0xFF:
        marker = data[off + 1]
        if marker == 0xDA:  # SOS: finiti i segmenti di metadati
            break
        if marker == 0xE1 and data[off + 4:off + 10] == b"Exif\x00\x00":
            return off + 10
        off += 2 + struct.unpack_from(">H", data, off + 2)[0]
    return None


def fix_broken_exif_ascii_tags(data):
    """Neutralizza, nel primo segmento Exif di un JPEG (anche se preceduto da
    altri segmenti, es. un JFIF), i tag ASCII (IFD0) il cui valore occupa
    4 byte o meno (es. Make='DJI').

    Bug di python-docx: per specifica TIFF un valore di 4 byte o meno va letto
    inline nel campo stesso, ma il parser Exif di python-docx lo tratta sempre
    come puntatore: legge quei 4 byte come un offset a vanvera e va in crash
    (UnexpectedEndOfFileError) durante add_picture(). Qui si ritipizza il tag
    da ASCII a BYTE (tipo TIFF valido, ma non gestito nello specifico da
    python-docx: viene ignorato senza leggere altro dallo stream). Cambia solo
    la copia in memoria usata per l'incorporazione, il file sorgente non è
    mai toccato; la firma Exif resta al suo posto così il formato è ancora
    riconosciuto correttamente.
    """
    tiff_start = _exif_tiff_start(data)
    if tiff_start is None:
        return data
    endian = data[tiff_start:tiff_start + 2]
    if endian not in (b"II", b"MM"):
        return data
    fmt = "<" if endian == b"II" else ">"
    out = bytearray(data)
    ifd0_offset = tiff_start + struct.unpack_from(fmt + "I", out, tiff_start + 4)[0]
    n_entries = struct.unpack_from(fmt + "H", out, ifd0_offset)[0]
    for idx in range(n_entries):
        entry_off = ifd0_offset + 2 + idx * 12
        field_type = struct.unpack_from(fmt + "H", out, entry_off + 2)[0]
        value_count = struct.unpack_from(fmt + "I", out, entry_off + 4)[0]
        if field_type == 2 and value_count <= 4:  # ASCII che andrebbe inline
            struct.pack_into(fmt + "H", out, entry_off + 2, 1)  # ritipizza a BYTE
    return bytes(out)


def _embed_picture(paragraph, path, width, an):
    """Prova a incorporare l'immagine; se fallisce ritenta con l'Exif corretto
    (vedi fix_broken_exif_ascii_tags). Ritorna True se incorporata, False altrimenti."""
    try:
        paragraph.add_run().add_picture(str(path), width=width)
        return True
    except Exception:
        pass
    try:
        raw = Path(path).read_bytes()
        fixed = fix_broken_exif_ascii_tags(raw)
        if fixed == raw:
            raise ValueError("no Exif tag to fix")
        paragraph.add_run().add_picture(io.BytesIO(fixed), width=width)
        if an is not None:
            an.add(Path(path).parent.name, "note: Exif fixed to embed the image",
                   "images", Path(path).name)
        return True
    except Exception as e:
        if an is not None:
            reason = "empty file (0 bytes)" if Path(path).stat().st_size == 0 else e.__class__.__name__
            an.add(Path(path).parent.name, "image cannot be embedded", "images",
                   f"{Path(path).name}: {reason}")
        return False


def compute_width(path, max_w, max_h):
    width = max_w
    size = image_size(path)
    if size and size[0] and size[1]:
        w, h = size
        width = min(max_w, int(max_h * w / h))
    return width


def add_image(container, path, cfg, caption=None, max_width=None, max_height=None,
              paragraph=None, an=None, space_before_pt=None, space_after_pt=None, align="center"):
    icfg = cfg.get("images", {})
    max_w = max_width if max_width is not None else Cm(icfg.get("max_width_cm", 15))
    max_h = max_height if max_height is not None else Cm(icfg.get("max_height_cm", 20))
    width = compute_width(path, max_w, max_h)
    p = paragraph if paragraph is not None else container.add_paragraph()
    p.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.CENTER)
    # Interlinea singola (auto) e nessun rientro: con l'interlinea "esatta" dello
    # stile Normal Word ritaglia l'immagine inline all'altezza di una riga.
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.first_line_indent = Pt(0)
    if space_before_pt is None:
        space_before_pt = icfg.get("figure_space_before_pt", 4)
    if space_after_pt is None:
        space_after_pt = icfg.get("figure_space_after_pt")   # None = stile Caption / 6pt senza didascalia
    p.paragraph_format.space_before = Pt(space_before_pt)
    # con didascalia lo spazio "sotto" va sulla didascalia stessa, non sull'immagine
    p.paragraph_format.space_after = Pt(0 if caption else (6 if space_after_pt is None else space_after_pt))
    p.paragraph_format.keep_with_next = bool(caption)
    if not _embed_picture(p, path, width, an):
        run = p.add_run(f"[image cannot be embedded: {Path(path).name}]")
        run.italic = True
    if caption:
        cp = container.add_paragraph(style="Caption")
        add_md_runs(cp, caption)
        if space_after_pt is not None:
            cp.paragraph_format.space_after = Pt(space_after_pt)


def add_bottom_border(pPr, size_eighths_pt, color, space_pt=2):
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_eighths_pt))
    bottom.set(qn("w:space"), str(space_pt))
    bottom.set(qn("w:color"), color)
    b.append(bottom)
    pPr.append(b)


def set_vertical_alignment(section, value):
    """value: 'top' | 'center' | 'both' | 'bottom' (allineamento verticale della sezione)."""
    sectPr = section._sectPr
    el = sectPr.find(qn("w:vAlign"))
    if el is None:
        el = OxmlElement("w:vAlign")
        sectPr.append(el)
    el.set(qn("w:val"), value)


def add_figures_title_page(doc, title):
    """Sezione dedicata a piena pagina con 'FIGURES' centrato in orizzontale e verticale,
    isolata da interruzioni di sezione prima e dopo."""
    title_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_vertical_alignment(title_section, "center")
    doc.add_heading(title, level=1)
    after_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_vertical_alignment(after_section, "top")


def add_quote(doc, q, cfg):
    """Citazione: rientrata a sinistra e destra (styles.quote). space_before/space_after
    dello stile sono lo spazio ESTERNO al blocco; fra i paragrafi interni c'e' paragraph_gap;
    la fonte segue a capo, allineata a destra, senza spazio sopra."""
    sc = style_cfg(cfg, "quote", fallback="normal")
    gap = sc.get("paragraph_gap", 4)
    paras = [x.strip() for x in re.split(r"\n\s*\n", q.text.strip()) if x.strip()]
    for i, para in enumerate(paras):
        p = doc.add_paragraph(style="Volumen Quote")
        add_md_runs(p, para)
        pf = p.paragraph_format
        pf.space_before = Pt(sc["space_before"] if i == 0 else 0)
        last = i == len(paras) - 1
        pf.space_after = Pt(0 if (last and q.source) else (sc["space_after"] if last else gap))
        pf.keep_with_next = bool(q.source) and last
        if last:
            pf.keep_together = True        # l'ultimo capoverso non si spezza fra due pagine
            pf.widow_control = True
    if q.source:
        p = doc.add_paragraph(style="Volumen Quote")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(sc["space_after"])
        add_md_runs(p, f"— {q.source}")


def add_list(doc, lb, cfg):
    """Elenco con marcatori e numeri scritti nel testo (niente numbering.xml): indentazione
    e sporgenza da styles.list (indent_left, first_line_indent negativo); ogni livello
    annidato aggiunge lo stesso rientro."""
    sc = style_cfg(cfg, "list", fallback="normal")
    step = sc["indent_left"] or 18
    counters = {}
    for depth, text in lb.items:
        for d in [k for k in counters if k > depth]:
            del counters[d]
        if lb.ordered:
            counters[depth] = counters.get(depth, 0) + 1
            marker = f"{counters[depth]}."
        else:
            marker = "•" if depth == 0 else "–"
        p = doc.add_paragraph(style="Volumen List")
        pf = p.paragraph_format
        pf.left_indent = Pt(step * (depth + 1))
        pf.tab_stops.add_tab_stop(Pt(step * (depth + 1)))
        p.add_run(f"{marker}\t")
        add_md_runs(p, text)


def add_part_title_page(doc, title, level, cfg, usable):
    """Pagina isolata con il solo titolo, centrato in orizzontale e verticale, senza
    intestazione corrente. La sezione successiva (capitolo seguente) la crea il capitolo."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_vertical_alignment(sec, "center")
    if cfg.get("header", {}).get("enabled", True):
        add_running_header(doc, sec, cfg, usable, title="", first=False)
    p = doc.add_heading(title, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def strip_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def add_portrait_pair(doc, cfg, usable, left, right, an=None):
    """Due immagini verticali affiancate in una tabella senza bordi."""
    gap = Cm(0.5)
    col_w = int((usable - gap) / 2)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    strip_table_borders(table)
    set_cell_margins(table)
    cells = table.rows[0].cells
    for cell, (path, caption) in zip(cells, (left, right)):
        cell.width = col_w
        add_image(cell, path, cfg, caption, max_width=col_w, paragraph=cell.paragraphs[0], an=an)


def add_figures_section(doc, cfg, usable, usable_h, figures, an=None):
    """Immagini orizzontali impilate N per foglio (piena larghezza); verticali a due
    affiancate quando adiacenti, altrimenti da sole ma alte quanto un'orizzontale
    (stesso "slot" di griglia), così una coppia orizzontale+verticale isolata
    entra sempre in un'unica pagina."""
    icfg = cfg.get("images", {})
    default_max_w = Cm(icfg.get("max_width_cm", 15))
    per_page = icfg.get("figures_landscape_per_page")
    landscape_max_h = Cm(icfg.get("max_height_cm", 20))     # senza figures_landscape_per_page: il tetto delle figure
    if per_page:
        # riserva per didascalia + spazio sopra/sotto ciascuna figura (figure_space_before/after_pt)
        reserve = (Cm(icfg.get("figures_row_reserve_cm", 1.2)) + Pt(icfg.get("figure_space_before_pt", 4))
                   + Pt(icfg.get("figure_space_after_pt") or 0))
        slot_h = int(usable_h / per_page)
        landscape_max_h = max(int(Cm(2)), slot_h - reserve)

    last_landscape_width = default_max_w  # riferimento finché non appare la prima orizzontale
    i, n = 0, len(figures)
    while i < n:
        path, caption = figures[i]
        if is_landscape(path):
            add_image(doc, path, cfg, caption, max_height=landscape_max_h, an=an)
            last_landscape_width = compute_width(path, default_max_w, landscape_max_h)
            i += 1
        elif i + 1 < n and not is_landscape(figures[i + 1][0]):
            add_portrait_pair(doc, cfg, usable, figures[i], figures[i + 1], an=an)
            i += 2
        else:
            # verticale isolata: stesso slot (altezza) e stessa larghezza dell'ultima orizzontale
            add_image(doc, path, cfg, caption, max_width=last_landscape_width,
                      max_height=landscape_max_h, an=an)
            i += 1


# --------------------------------------------------------------------------
# Volume -> documento
# --------------------------------------------------------------------------
def add_toc_field(doc, depth):
    """Campo indice: Word lo popola all'apertura (accetta l'aggiornamento dei campi)
    o con tasto destro > Aggiorna campo."""
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, text in (("begin", None), (None, f'TOC \\o "1-{depth}" \\h \\z \\u'),
                       ("separate", None), (None, "Aggiornare il campo per generare l'indice."),
                       ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                el.set(qn("w:dirty"), "true")   # chiede a Word di aggiornarlo all'apertura
            run._r.append(el)
        elif text.startswith("TOC"):
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
    doc.add_page_break()


DOC = None               # documento in costruzione (serve ai commenti, che stanno nel documento principale)


def build_document(cfg, volume, an):
    global HIGHLIGHT
    HIGHLIGHT = cfg.get("typography", {}).get("highlight_color", HIGHLIGHT)
    global LINKS
    LINKS = cfg.get("links", {})
    doc, usable, usable_h = setup_document(cfg)
    global DOC
    DOC = doc
    icfg = cfg.get("images", {})
    fig_classes = icfg.get("classes", {})
    table_classes = cfg.get("table", {}).get("classes", {})

    first_sec = doc.sections[0]
    if volume.front:
        # frontespizio e colophon: prima sezione senza intestazione ne' numero; il resto parte su pagina destra
        # e la numerazione riparte (come nel PDF)
        for fi, fch in enumerate(volume.front):
            if fi > 0:
                doc.add_page_break()
            for b in fch.blocks:
                if isinstance(b, model.StyledText):
                    add_styled_text(doc, b, usable / 360000)
                elif isinstance(b, model.PageBreak):
                    doc.add_page_break()
                elif isinstance(b, model.Text):
                    add_text_block(doc, b.text)
                elif isinstance(b, model.Heading):
                    doc.add_heading(b.text, level=b.level)
                elif isinstance(b, model.Figure):
                    usable_cm = usable / 360000
                    w_cm = model.length_cm(b.width, usable_cm)
                    h_cm = model.length_cm(b.height, usable_cm)
                    add_image(doc, b.path, cfg, "", an=an, max_width=Cm(w_cm) if w_cm else Cm(1000),
                              max_height=Cm(h_cm) if h_cm else Cm(1000), align=b.align or "center",
                              space_before_pt=0, space_after_pt=0)
        first_sec = doc.add_section(WD_SECTION.ODD_PAGE)
        set_vertical_alignment(first_sec, "top")

    toc = volume.meta.get("toc") or {}
    if toc.get("enabled"):
        add_toc_field(doc, toc.get("depth", 2))

    hcfg = cfg.get("header", {})
    bookmark_id = [0]
    for i, ch in enumerate(volume.chapters):
        starts_with_part = bool(ch.blocks) and isinstance(ch.blocks[0], model.PartTitle)
        running = ch.running_title or ch.title or hcfg.get("title", "")
        if hcfg.get("enabled", True) and not starts_with_part:
            if i == 0:
                add_running_header(doc, first_sec, cfg, usable, title=running, first=True,
                                   left_title=model.running_left(ch, volume, hcfg))
            elif ch.new_page:
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
                set_vertical_alignment(sec, "top")
                add_running_header(doc, sec, cfg, usable, title=running, first=False,
                                   left_title=model.running_left(ch, volume, hcfg))
        elif i > 0 and ch.new_page and not starts_with_part:
            doc.add_section(WD_SECTION.NEW_PAGE)
        for b in ch.blocks:
            if isinstance(b, model.Heading):
                h = doc.add_heading(b.text, level=b.level)
                first = bool(ch.blocks) and ch.blocks[0] is b
                if first:
                    bookmark_id[0] += 1
                    add_bookmark(h, model.chapter_anchor(ch.id), bookmark_id[0])
                elif b.anchor:                      # scheda di un complesso: destinazione dei link [ID](ID.md)
                    bookmark_id[0] += 1
                    add_bookmark(h, model.chapter_anchor(b.anchor), bookmark_id[0])
                if b.level == 1 and first:
                    # titolo h1 di capitolo: parte a offset_pct dell'altezza utile
                    h.paragraph_format.space_before = Pt(usable_h / 12700
                                                        * cfg.get("chapter_opener", {}).get("offset_pct", 33) / 100)
                if first and ch.authors:
                    # sotto il titolo (di qualsiasi livello: anche i sotto-capitoli), gli autori in grigio
                    ap = doc.add_paragraph(style="Volumen Chapter Author")
                    ap.add_run(", ".join(ch.authors))
                if b.page_break_before:
                    h.paragraph_format.page_break_before = True
            elif isinstance(b, model.PartTitle):
                add_part_title_page(doc, b.text, b.level, cfg, usable)
            elif isinstance(b, model.Text):
                add_text_block(doc, b.text)
            elif isinstance(b, model.StyledText):
                add_styled_text(doc, b, usable / 360000)
            elif isinstance(b, model.Bibliography):
                doc.add_heading(b.title, level=b.level)
                for _, entry in b.entries:
                    add_md_runs(doc.add_paragraph(style="Volumen Bibliography"), entry)
            elif isinstance(b, model.Quote):
                add_quote(doc, b, cfg)
            elif isinstance(b, model.ListBlock):
                add_list(doc, b, cfg)
            elif isinstance(b, model.Figure) and b.klass in fig_classes:
                # Classe di figura (images.classes.<nome>): larghezza scollegata da images.max_width_cm/
                # max_height_cm (width_pct della larghezza utile), altezza secondo le proporzioni.
                fc = fig_classes[b.klass]
                add_image(doc, b.path, cfg, b.caption, an=an,
                          max_width=int(usable * fc.get("width_pct", 100) / 100), max_height=Cm(1000),
                          space_before_pt=fc.get("space_before_pt", icfg.get("figure_space_before_pt", 4)),
                          space_after_pt=fc.get("space_after_pt", icfg.get("figure_space_before_pt", 4)))
            elif isinstance(b, model.Figure):
                # width/height espliciti (attributi Markdown): scavalcano i tetti; % = larghezza utile
                usable_cm = usable / 360000
                w_cm = model.length_cm(b.width, usable_cm)
                h_cm = model.length_cm(b.height, usable_cm)
                kw = {}
                if w_cm or h_cm:
                    kw = {"max_width": Cm(w_cm) if w_cm else Cm(1000),
                          "max_height": Cm(h_cm) if h_cm else Cm(1000)}
                add_image(doc, b.path, cfg, b.caption, an=an, **kw)
            elif isinstance(b, model.Table):
                tc = table_classes.get(b.klass, {})      # classe di tabella: table.classes.<nome>
                add_table(doc, b.header, b.rows, cfg, usable, ratios=tc.get("col_ratios") or b.ratios, aligns=b.aligns,
                          caption=b.caption, empty=tc.get("empty_cell", cfg.get("table", {}).get("empty_cell", "–")),
                          markup=tc.get("markup", True))
            elif isinstance(b, model.FigureGroup):
                if b.title is not None:
                    add_figures_title_page(doc, b.title)
                add_figures_section(doc, cfg, usable, usable_h,
                                    [(f.path, f.caption) for f in b.figures], an=an)
            else:
                raise TypeError(f"blocco non gestito dal renderer DOCX: {b!r}")
    return doc
